"""Build a queryable SQLite index from PSP's spreadsheets and CV folder.

Three sources, three tables:

  candidates   MAIN CANDIDATE LIST.xlsx  - one tab per region, contact + location
  submissions  Conor Submission Spreadsheet.xlsx - one tab per client, full history
  cvs          Candidate CVs/            - one .docx per candidate, specialism evidence

Names are the join key across all three, so they are normalised hard: case,
punctuation and trailing whitespace all vary between sources.
"""
from __future__ import annotations

import argparse
import os
import re
import sqlite3
import zipfile
from datetime import datetime, date

from . import domain, geo

SCHEMA = """
CREATE TABLE IF NOT EXISTS candidates (
    name_key    TEXT PRIMARY KEY,
    name        TEXT,
    address     TEXT,
    postcode    TEXT,
    outward     TEXT,
    phone       TEXT,
    email       TEXT,
    sector      TEXT,
    position    TEXT,
    grade       TEXT,
    notes       TEXT,
    region_tab  TEXT
);
CREATE TABLE IF NOT EXISTS submissions (
    name_key    TEXT,
    name        TEXT,
    client      TEXT,
    submitted   TEXT,
    job_title   TEXT,
    outcome_raw TEXT,
    outcome     TEXT,
    grade       TEXT,
    specialisms TEXT
);
CREATE TABLE IF NOT EXISTS cvs (
    name_key    TEXT PRIMARY KEY,
    name        TEXT,
    path        TEXT,
    modified    TEXT,
    specialisms TEXT,
    excerpt     TEXT
);
CREATE INDEX IF NOT EXISTS ix_sub_name ON submissions(name_key);
CREATE INDEX IF NOT EXISTS ix_sub_client ON submissions(client);
CREATE INDEX IF NOT EXISTS ix_cand_outward ON candidates(outward);
"""


def name_key(name: str | None) -> str:
    """Normalise a person's name into a stable join key."""
    if not name:
        return ""
    n = re.sub(r"\([^)]*\)", " ", str(name))          # drop "(Bexley)" style notes
    n = re.sub(r"[^A-Za-z' -]+", " ", n).lower()
    n = re.sub(r"\s+", " ", n).strip()
    return n


# The outcome column is free text typed over two years. These patterns are
# ordered so that the more specific state wins: "interview but didn't go" is a
# withdrawal, not an interview we should count as progress.
OUTCOME_PATTERNS: list[tuple[str, str]] = [
    (r"offer", "offered"),
    (r"another agency|went directly|applied directly|new job|new role|got a job",
     "lost_to_other"),
    # normalise() strips curly apostrophes to spaces, so "didn't" arrives as
    # "didn t" - the gap tolerance below covers every variant typed so far.
    (r"didn.{0,3}t go|didnr go|never went|couldn.{0,3}t go|not going|but didn"
     r"|withdrew|ignored us|too far|no longer available|cancelled",
     "interview_withdrawn"),
    (r"interview|\biv\b|test|panel", "interviewed"),
    (r"role closed|on hold|no roles|no sponsorship", "role_dead"),
]


def classify_outcome(raw: str | None) -> str:
    """Reduce the free-text interview column to a comparable state."""
    t = domain.normalise(raw)
    if not t.strip():
        return "submitted"
    for pattern, state in OUTCOME_PATTERNS:
        if re.search(pattern, t):
            return state
    if re.match(r"^\s*\d{1,2}[/-]\d{1,2}[/-]\d{2,4}", str(raw).strip()):
        return "interviewed"
    return "submitted"


def _as_date(value) -> str | None:
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    if not value:
        return None
    for fmt in ("%m/%d/%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(str(value).strip(), fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def _headers(row) -> dict[str, int]:
    return {domain.normalise(c).strip(): i for i, c in enumerate(row) if c}


def _pick(hdr: dict[str, int], *names: str) -> int | None:
    for n in names:
        for key, idx in hdr.items():
            if key.startswith(n):
                return idx
    return None


def load_candidates(conn: sqlite3.Connection, path: str) -> int:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    rows, seen = [], set()
    for sheet in wb.worksheets:
        it = sheet.iter_rows(values_only=True)
        try:
            hdr = _headers(next(it))
        except StopIteration:
            continue
        i_name = _pick(hdr, "name")
        if i_name is None:
            continue
        i_addr, i_pc = _pick(hdr, "address"), _pick(hdr, "postcode", "post code")
        i_num, i_email = _pick(hdr, "number", "phone"), _pick(hdr, "email")
        i_sector = _pick(hdr, "adults/childrens", "adults", "sector")
        i_pos, i_notes = _pick(hdr, "current position", "position"), _pick(hdr, "notes")

        def cell(row, idx):
            if idx is None or idx >= len(row):
                return None
            v = row[idx]
            return str(v).strip() if v not in (None, "") else None

        for row in it:
            name = cell(row, i_name)
            if not name:
                continue
            key = name_key(name)
            # Regional tabs overlap heavily; first sighting wins.
            if not key or key in seen:
                continue
            seen.add(key)
            postcode = cell(row, i_pc)
            position = cell(row, i_pos)
            rows.append((
                key, name, cell(row, i_addr), postcode, geo.outward(postcode),
                cell(row, i_num), cell(row, i_email),
                domain.parse_sector(cell(row, i_sector)) or domain.parse_sector(position),
                position, domain.parse_grade(position), cell(row, i_notes), sheet.title,
            ))
    wb.close()
    conn.executemany(
        "INSERT OR REPLACE INTO candidates VALUES (?,?,?,?,?,?,?,?,?,?,?,?)", rows)
    return len(rows)


def load_submissions(conn: sqlite3.Connection, path: str) -> int:
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True, data_only=True)
    rows = []
    for sheet in wb.worksheets:
        client = sheet.title.strip()
        it = sheet.iter_rows(values_only=True)
        try:
            hdr = _headers(next(it))
        except StopIteration:
            continue
        i_name = _pick(hdr, "name")
        i_date = _pick(hdr, "date submitted", "date")
        i_job = _pick(hdr, "job title", "job", "role")
        i_out = _pick(hdr, "interview date", "interview", "outcome", "status")
        if i_name is None or i_job is None:
            continue
        for row in it:
            name = row[i_name] if i_name < len(row) else None
            if not name or not str(name).strip():
                continue
            job = str(row[i_job]).strip() if i_job < len(row) and row[i_job] else ""
            raw = str(row[i_out]).strip() if i_out is not None and i_out < len(row) and row[i_out] else ""
            rows.append((
                name_key(name), str(name).strip(), client,
                _as_date(row[i_date] if i_date is not None and i_date < len(row) else None),
                job, raw, classify_outcome(raw),
                domain.parse_grade(job),
                ",".join(sorted(domain.parse_specialisms(job))),
            ))
    wb.close()
    conn.executemany(
        "INSERT INTO submissions VALUES (?,?,?,?,?,?,?,?,?)", rows)
    return len(rows)


DOCX_TEXT = re.compile(r"<w:t[^>]*>([^<]*)</w:t>")


def read_docx(path: str, limit: int = 20000) -> str:
    """Pull plain text out of a .docx without a dependency on python-docx."""
    try:
        with zipfile.ZipFile(path) as z:
            xml = z.read("word/document.xml").decode("utf-8", "ignore")
    except (zipfile.BadZipFile, KeyError, OSError):
        return ""
    return re.sub(r"\s+", " ", " ".join(DOCX_TEXT.findall(xml))[:limit]).strip()


def load_cvs(conn: sqlite3.Connection, folder: str) -> int:
    rows = []
    for entry in sorted(os.scandir(folder), key=lambda e: e.name):
        if not entry.is_file() or not entry.name.lower().endswith(".docx"):
            continue
        if entry.name.startswith("~$"):
            continue
        person = re.sub(r"\s*\bcv\b.*$", "", os.path.splitext(entry.name)[0],
                        flags=re.IGNORECASE)
        text = read_docx(entry.path)
        rows.append((
            name_key(person), person.strip(), entry.path,
            datetime.fromtimestamp(entry.stat().st_mtime).strftime("%Y-%m-%d"),
            ",".join(sorted(domain.parse_specialisms(text))), text[:1200],
        ))
    conn.executemany("INSERT OR REPLACE INTO cvs VALUES (?,?,?,?,?,?)", rows)
    return len(rows)


def build(db_path: str, master: str | None, subs: str | None, cvs: str | None) -> dict:
    conn = sqlite3.connect(db_path)
    conn.executescript(SCHEMA)
    counts = {}
    if master:
        counts["candidates"] = load_candidates(conn, master)
    if subs:
        conn.execute("DELETE FROM submissions")
        counts["submissions"] = load_submissions(conn, subs)
    if cvs:
        counts["cvs"] = load_cvs(conn, cvs)
    conn.commit()
    conn.close()
    return counts


def main() -> None:
    p = argparse.ArgumentParser(description="Build the PSP candidate index.")
    p.add_argument("--db", default="data/psp.db")
    p.add_argument("--master", help="MAIN CANDIDATE LIST.xlsx")
    p.add_argument("--subs", help="Conor Submission Spreadsheet.xlsx")
    p.add_argument("--cvs", help="Candidate CVs folder")
    a = p.parse_args()
    os.makedirs(os.path.dirname(os.path.abspath(a.db)), exist_ok=True)
    for k, v in build(a.db, a.master, a.subs, a.cvs).items():
        print(f"{k}: {v:,} rows")


if __name__ == "__main__":
    main()
